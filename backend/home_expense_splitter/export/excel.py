# home_expense_splitter/export/excel.py
"""Genera il .xlsx del divisore: il foglio Riepilogo Finale.

Il saldo è calcolato con lo stesso motore di `split`, sulla sorgente configurata. I fogli
di dettaglio (spese condivise/personali, bonifici) sono del libro mastro: li produce
l'export di `finance` (`GET /api/export.xlsx`), non questo nodo (issue #4).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from home_expense_splitter.config import load_split_config, member_names
from home_expense_splitter.split.engine import compute_split

EURO = '#,##0.00\\ "€"'
_TITLE = Font(bold=True, size=13, color="FFFFFF")
_TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="2E75B6")
_BOLD = Font(bold=True)


def _title(ws: Worksheet, text: str, span: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(1, 1, text)
    c.font, c.fill, c.alignment = _TITLE, _TITLE_FILL, Alignment(horizontal="center")


def _header(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, name in enumerate(headers, start=1):
        c = ws.cell(row, col, name)
        c.font, c.fill = _HEAD, _HEAD_FILL


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(session: Session) -> Workbook:
    cfg = load_split_config()
    name_a, name_b = member_names(cfg)
    res = compute_split(
        session,
        member_a_pct=cfg["percentuali"]["member_a"],
        member_b_pct=cfg["percentuali"]["member_b"],
        escluse=cfg["escluse"],
    )

    wb = Workbook()

    # --- Riepilogo Finale ---
    ws = wb.active
    ws.title = "Riepilogo Finale"
    _title(ws, "RIEPILOGO DIVISIONE SPESE", 2)
    _header(ws, 3, ["Categoria condivisa (divisibile)", "Totale"])
    r = 4
    for cat, cents in res["per_category"]:
        ws.cell(r, 1, cat or "(senza categoria)")
        ws.cell(r, 2, cents / 100).number_format = EURO
        r += 1

    def kv(label: str, cents: int, bold: bool = False) -> None:
        nonlocal r
        a = ws.cell(r, 1, label)
        b = ws.cell(r, 2, cents / 100)
        b.number_format = EURO
        if bold:
            a.font = b.font = _BOLD
        r += 1

    r += 1
    kv("TOTALE DIVISIBILE", res["divisible_cents"], bold=True)
    kv(f"Quota {name_b} ({res['member_b_pct']}%)", res["member_b_quota_cents"])
    kv(f"Bonifici già dati da {name_b} (netto)", res["net_bonifici_cents"])
    kv(f"{name_b.upper()} DEVE A {name_a.upper()}", res["saldo_cents"], bold=True)
    _autosize(ws, [42, 16])

    return wb


def export_workbook(session: Session, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(session).save(path)
    return path
